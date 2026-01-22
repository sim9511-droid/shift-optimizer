import streamlit as st
import pandas as pd
import datetime
from shift_optimizer import solve
import os

st.set_page_config(page_title="シフト最適化システム", page_icon="📅", layout="wide")

# 定数定義
WEEKDAY_MAP = {
    'Monday': '月', 'Tuesday': '火', 'Wednesday': '水', 
    'Thursday': '木', 'Friday': '金', 'Saturday': '土', 'Sunday': '日'
}

# ヘルパー関数
def clean_dataframe(df):
    """DataFrameをArrow互換の文字列型に変換"""
    return df.astype(str)

def remove_none_values(df, exclude_cols=[]):
    """DataFrameからNone/NaN値を空文字列に置換"""
    for col in df.columns:
        if col not in exclude_cols:
            df[col] = df[col].apply(
                lambda x: '' if (x is None or pd.isna(x) or str(x).strip().lower() in ['none', 'nan', '']) 
                else str(x).strip()
            )
    return df

st.title("📅 シフト最適化システム")
st.markdown("---")

# 現在のフォルダパスを取得
current_dir = os.getcwd()

# サイドバーで入力ファイルの確認
with st.sidebar:
    st.header("ℹ️ システム情報")
    st.info(f"""
    このアプリは、スタッフのシフトを自動最適化します。
    
    **作業フォルダ:**
    `{current_dir}`
    
    **必要なファイル:**
    - Shift_Input.xlsx
    - constraints.yaml
    """)
    
    # ファイルの存在確認
    input_file_exists = os.path.exists('Shift_Input.xlsx')
    constraints_file_exists = os.path.exists('constraints.yaml')
    
    st.subheader("📁 ファイルチェック")
    if input_file_exists:
        st.success("✅ Shift_Input.xlsx")
        # ファイル情報を表示
        file_stat = os.stat('Shift_Input.xlsx')
        file_time = datetime.datetime.fromtimestamp(file_stat.st_mtime)
        st.caption(f"更新日時: {file_time.strftime('%Y/%m/%d %H:%M')}")
    else:
        st.error("❌ Shift_Input.xlsx が見つかりません")
    
    if constraints_file_exists:
        st.success("✅ constraints.yaml")
        # ファイル情報を表示
        file_stat = os.stat('constraints.yaml')
        file_time = datetime.datetime.fromtimestamp(file_stat.st_mtime)
        st.caption(f"更新日時: {file_time.strftime('%Y/%m/%d %H:%M')}")
    else:
        st.error("❌ constraints.yaml が見つかりません")

# タブを作成
tab1, tab2, tab3, tab4 = st.tabs(["🚀 シフト最適化", "📋 設定ファイル確認", "📝 希望休入力", "⚙️ 制約条件設定"])

# タブ1: シフト最適化
with tab1:
    st.subheader("🚀 シフト最適化を実行")
    st.write("下のボタンをクリックして、シフトの最適化を開始します。")
    
    if st.button("▶️ シフトを最適化する", type="primary", use_container_width=True):
        if not input_file_exists or not constraints_file_exists:
            st.error("⚠️ 必要なファイルが見つかりません。ファイルを確認してください。")
        else:
            with st.spinner("🔄 シフトを最適化中... しばらくお待ちください..."):
                try:
                    # シフト最適化を実行
                    solve()
                    st.success("✅ シフトの最適化が完了しました！")
                    st.balloons()
                    
                    # 結果ファイルの確認
                    if os.path.exists('shift_schedule_final.xlsx'):
                        st.info("📄 結果ファイル: `shift_schedule_final.xlsx` が作成されました。")
                        
                        # 結果をプレビュー
                        try:
                            df_result = pd.read_excel('shift_schedule_final.xlsx', sheet_name='シフト表')
                            st.subheader("📊 シフト表プレビュー")
                            # データ型を文字列に変換してArrowエラーを回避
                            df_result = df_result.astype(str)
                            st.dataframe(df_result, use_container_width=True, height=400)
                            
                            # ダウンロードボタン
                            with open('shift_schedule_final.xlsx', 'rb') as f:
                                st.download_button(
                                    label="⬇️ シフト表をダウンロード",
                                    data=f,
                                    file_name=f"shift_schedule_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
                                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                    use_container_width=True
                                )
                        except Exception as e:
                            st.warning(f"プレビューの表示に失敗しました: {e}")
                    
                except Exception as e:
                    st.error(f"❌ エラーが発生しました: {str(e)}")
                    st.exception(e)
    
    # 使い方
    with st.expander("📝 使い方"):
        st.markdown("""
        1. **Shift_Input.xlsx** を準備
           - 「設定」シート: 年月と人数設定
           - 「休日回数」シート: スタッフ情報
           - 「希望休」シート: 休日希望
        
        2. **constraints.yaml** を設定
           - シフト期間
           - 祝日リスト
           - 制約条件
        
        3. 「シフトを最適化する」ボタンをクリック
        
        4. 結果をダウンロード
        """)

# タブ2: 設定ファイル確認
with tab2:
    st.subheader("📋 設定ファイルの内容確認")
    
    if os.path.exists('Shift_Input.xlsx'):
        st.markdown("### 📊 Shift_Input.xlsx")
        
        # 各シートを読み込んで表示
        excel_file = pd.ExcelFile('Shift_Input.xlsx')
        sheet_names = excel_file.sheet_names
        
        selected_sheet = st.selectbox("表示するシートを選択", sheet_names)
        
        if selected_sheet == '設定':
            df = pd.read_excel('Shift_Input.xlsx', sheet_name=selected_sheet, header=None)
            st.dataframe(clean_dataframe(df), use_container_width=True)
            st.info("""
            **設定シートの構造:**
            - A1: 年, B1: 値
            - A2: 月, B2: 値
            - A4: 平日最低人数, B4: 値
            - A5: 平日最大人数, B5: 値
            - A6: 休日最低人数, B6: 値
            - A7: 休日最大人数, B7: 値
            - D列: 祝日リスト（日付形式）
            """)
        elif selected_sheet in ['休日回数', '希望休']:
            df = pd.read_excel('Shift_Input.xlsx', sheet_name=selected_sheet, header=0)
            st.dataframe(clean_dataframe(df), use_container_width=True)
            info_msg = "スタッフの基本情報、月間休日数、勤務時間が設定されています。" if selected_sheet == '休日回数' else "各スタッフの希望休が「希」「有」「休」で記載されています。"
            st.info(info_msg)
        else:
            df = pd.read_excel('Shift_Input.xlsx', sheet_name=selected_sheet)
            st.dataframe(clean_dataframe(df), use_container_width=True)
    else:
        st.warning("⚠️ Shift_Input.xlsx が見つかりません")
    
    st.markdown("---")
    
    if os.path.exists('constraints.yaml'):
        st.markdown("### 📄 constraints.yaml")
        with open('constraints.yaml', 'r', encoding='utf-8') as f:
            yaml_content = f.read()
        st.code(yaml_content, language='yaml')
    else:
        st.warning("⚠️ constraints.yaml が見つかりません")

# タブ3: 希望休入力
with tab3:
    st.subheader("📝 希望休の入力・編集")
    
    if os.path.exists('Shift_Input.xlsx'):
        try:
            # 希望休シートを読み込み
            df_wish = pd.read_excel('Shift_Input.xlsx', sheet_name='希望休', header=0)
            
            # 全データをまず文字列型に変換してからNoneを処理
            for col in df_wish.columns:
                if col != '日付':
                    # 明示的に文字列に変換
                    df_wish[col] = df_wish[col].astype('object')
                    df_wish[col] = df_wish[col].apply(lambda x: '' if (pd.isna(x) or x is None or str(x).lower() == 'none' or str(x).lower() == 'nan') else str(x) if x != '' else '')
            
            # 設定シートから祝日リストを取得
            df_settings = pd.read_excel('Shift_Input.xlsx', sheet_name='設定', header=None)
            holidays = []
            for i in range(len(df_settings)):
                if pd.notna(df_settings.iloc[i, 3]):  # D列（index=3）
                    try:
                        holiday_date = pd.to_datetime(df_settings.iloc[i, 3])
                        holidays.append(holiday_date.date())
                    except:
                        pass
            
            # 日付列を日付型に変換（時間部分を削除）
            if '日付' in df_wish.columns:
                df_wish['日付'] = pd.to_datetime(df_wish['日付']).dt.date
            
            # 曜日列を追加（絵文字で視覚的に区別）
            df_wish_display = df_wish.copy()
            if '日付' in df_wish_display.columns:
                # 曜日を計算
                weekday_series = pd.to_datetime(df_wish_display['日付']).dt.day_name()
                
                # 曜日ラベルを作成（祝日、土日を絵文字で区別）
                weekday_labels = []
                for i, date_val in enumerate(df_wish_display['日付']):
                    weekday = weekday_series.iloc[i]
                    weekday_jp = {'Monday': '月', 'Tuesday': '火', 'Wednesday': '水', 
                                 'Thursday': '木', 'Friday': '金', 'Saturday': '土', 'Sunday': '日'}[weekday]
                    
                    # 祝日チェック
                    if date_val in holidays:
                        weekday_labels.append(f'🔴 {weekday_jp}')  # 祝日
                    elif weekday == 'Saturday':
                        weekday_labels.append(f'🔵 {weekday_jp}')  # 土曜日
                    elif weekday == 'Sunday':
                        weekday_labels.append(f'🔴 {weekday_jp}')  # 日曜日
                    else:
                        weekday_labels.append(weekday_jp)
                
                df_wish_display['曜日'] = weekday_labels
                
                # 曜日列を日付の次に配置
                cols = list(df_wish_display.columns)
                date_idx = cols.index('日付')
                cols.insert(date_idx + 1, cols.pop(cols.index('曜日')))
                df_wish_display = df_wish_display[cols]
            
            # スタッフ列のNone/NaNを空文字列に置換（日付・曜日列以外）
            for col in df_wish_display.columns:
                if col not in ['日付', '曜日', 'Unnamed: 0']:
                    # 二重チェック：確実に空文字列にする
                    df_wish_display[col] = df_wish_display[col].apply(
                        lambda x: '' if (x is None or pd.isna(x) or str(x).strip().lower() in ['none', 'nan', '']) else str(x).strip()
                    )
            
            st.info("各スタッフの希望休を入力してください。空白、希休、有休から選択できます。")
            
            # データエディタで編集可能にする
            column_config = {}
            
            # 日付列の設定
            if '日付' in df_wish_display.columns:
                column_config['日付'] = st.column_config.DateColumn(
                    '日付',
                    format="YYYY/MM/DD",
                    disabled=True
                )
            
            # 曜日列の設定（編集不可、色付き表示）
            if '曜日' in df_wish_display.columns:
                column_config['曜日'] = st.column_config.TextColumn(
                    '曜日',
                    disabled=True,
                    width="small"
                )
            
            # Unnamed列は非表示
            if 'Unnamed: 0' in df_wish_display.columns:
                column_config['Unnamed: 0'] = None
            
            # スタッフ列をプルダウン対応にする
            for col in df_wish_display.columns:
                if col not in ['日付', '曜日', 'Unnamed: 0']:
                    column_config[col] = st.column_config.SelectboxColumn(
                        col,
                        options=["", "希休", "有休"],
                        required=False
                    )
            
            # 曜日に応じたスタイリング用のCSS（罫線を強化）
            st.markdown("""
            <style>
            /* データエディタの罫線を強化 */
            div[data-testid="stDataEditor"] div[data-testid="stDataFrameResizable"] > div {
                border: 2px solid #666 !important;
            }
            div[data-testid="stDataEditor"] table {
                border-collapse: collapse !important;
            }
            div[data-testid="stDataEditor"] th {
                border: 1px solid #999 !important;
                border-bottom: 2px solid #666 !important;
                padding: 8px !important;
                background-color: #f0f0f0 !important;
            }
            div[data-testid="stDataEditor"] td {
                border: 1px solid #ccc !important;
                padding: 6px !important;
            }
            div[data-testid="stDataEditor"] tbody tr {
                border-bottom: 1px solid #ddd !important;
            }
            /* セルのエディタ部分にも罫線 */
            div[data-testid="stDataEditor"] input,
            div[data-testid="stDataEditor"] select {
                border: 1px solid #ccc !important;
            }
            </style>
            """, unsafe_allow_html=True)
            
            # 編集可能なデータフレーム
            edited_df = st.data_editor(
                df_wish_display,
                column_config=column_config,
                num_rows="fixed",
                use_container_width=True,
                height=500,
                hide_index=True
            )
            
            # 保存ボタン
            col1, col2 = st.columns(2)
            with col1:
                if st.button("💾 希望休を保存", type="primary"):
                    try:
                        from openpyxl import load_workbook
                        from openpyxl.styles import PatternFill, Font, Border, Side
                        
                        # 既存のExcelファイルを読み込み
                        with pd.ExcelFile('Shift_Input.xlsx') as xls:
                            # 全シートを読み込み
                            sheets = {}
                            for sheet_name in xls.sheet_names:
                                if sheet_name != '希望休':
                                    sheets[sheet_name] = pd.read_excel(xls, sheet_name=sheet_name, header=None if sheet_name == '設定' else 0)
                        
                        # 曜日列を削除して元の形式に戻す
                        df_to_save = edited_df.copy()
                        if '曜日' in df_to_save.columns:
                            df_to_save = df_to_save.drop(columns=['曜日'])
                        
                        # 空文字列をNaNに戻す（元のExcel形式に合わせる）
                        df_to_save = df_to_save.replace('', None)
                        
                        # 希望休シートを更新
                        sheets['希望休'] = df_to_save
                        
                        # Excelファイルに書き戻し
                        with pd.ExcelWriter('Shift_Input.xlsx', engine='openpyxl') as writer:
                            for sheet_name, df in sheets.items():
                                df.to_excel(writer, sheet_name=sheet_name, index=False, header=True if sheet_name != '設定' else False)
                        
                        # スタイリングを適用
                        wb = load_workbook('Shift_Input.xlsx')
                        ws = wb['希望休']
                        
                        # 罫線の定義
                        thin_border = Border(
                            left=Side(style='thin'),
                            right=Side(style='thin'),
                            top=Side(style='thin'),
                            bottom=Side(style='thin')
                        )
                        
                        # 曜日列を追加してスタイリング
                        ws.insert_cols(2)  # B列に挿入
                        ws.cell(1, 2, '曜日')  # ヘッダー
                        
                        # 各行に曜日を追加してスタイリング
                        for row_idx in range(2, ws.max_row + 1):
                            date_cell = ws.cell(row_idx, 1)
                            if date_cell.value:
                                try:
                                    date_val = pd.to_datetime(date_cell.value).date()
                                    weekday = pd.to_datetime(date_val).day_name()
                                    weekday_jp = {'Monday': '月', 'Tuesday': '火', 'Wednesday': '水', 
                                                 'Thursday': '木', 'Friday': '金', 'Saturday': '土', 'Sunday': '日'}[weekday]
                                    
                                    weekday_cell = ws.cell(row_idx, 2, weekday_jp)
                                    
                                    # 土日祝の色分け
                                    if date_val in holidays:
                                        # 祝日：背景赤、文字白、太字
                                        weekday_cell.fill = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')
                                        weekday_cell.font = Font(color='FFFFFF', bold=True)
                                    elif weekday == 'Saturday':
                                        # 土曜日：青文字
                                        weekday_cell.font = Font(color='0000FF')
                                    elif weekday == 'Sunday':
                                        # 日曜日：赤文字
                                        weekday_cell.font = Font(color='FF0000')
                                    
                                    weekday_cell.border = thin_border
                                except:
                                    pass
                        
                        # 全セルに罫線を適用
                        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
                            for cell in row:
                                cell.border = thin_border
                        
                        wb.save('Shift_Input.xlsx')
                        
                        st.success("✅ 希望休を保存しました！（曜日の色分けも適用されました）")
                        st.balloons()
                        st.rerun()
                    except Exception as e:
                        st.error(f"❌ 保存に失敗しました: {e}")
                        import traceback
                        st.code(traceback.format_exc())
            
            with col2:
                # ダウンロードボタン
                with open('Shift_Input.xlsx', 'rb') as f:
                    st.download_button(
                        label="⬇️ 更新したファイルをダウンロード",
                        data=f,
                        file_name="Shift_Input.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
        
        except Exception as e:
            st.error(f"❌ エラーが発生しました: {e}")
            st.exception(e)
    else:
        st.warning("⚠️ Shift_Input.xlsx が見つかりません")

# タブ4: 制約条件設定
with tab4:
    st.subheader("⚙️ 制約条件の編集")
    
    if os.path.exists('constraints.yaml'):
        import yaml
        
        with open('constraints.yaml', 'r', encoding='utf-8') as f:
            constraints = yaml.safe_load(f)
        
        st.markdown("### 📝 YAML設定の編集")
        
        # YAMLをテキストエリアで編集可能にする
        edited_yaml = st.text_area(
            "constraints.yaml の内容を編集",
            value=open('constraints.yaml', 'r', encoding='utf-8').read(),
            height=400
        )
        
        col1, col2 = st.columns(2)
        with col1:
            if st.button("💾 保存", type="primary"):
                try:
                    # YAML形式が正しいか検証
                    yaml.safe_load(edited_yaml)
                    with open('constraints.yaml', 'w', encoding='utf-8') as f:
                        f.write(edited_yaml)
                    st.success("✅ constraints.yaml を保存しました")
                    st.rerun()
                except yaml.YAMLError as e:
                    st.error(f"❌ YAML形式エラー: {e}")
        
        with col2:
            if st.button("🔄 元に戻す"):
                st.rerun()
        
        st.markdown("---")
        st.info("""
        **設定項目の説明:**
        - `period.start_day`: シフト開始日
        - `period.end_day`: シフト終了日
        - `period.holidays`: 祝日リスト（YYYY-MM-DD形式）※ 現在は設定シートD列から読み込みます
        - `rules.max_consecutive_work_days`: 最大連続勤務日数
        """)
    else:
        st.warning("⚠️ constraints.yaml が見つかりません")

# フッター
st.markdown("---")
st.markdown(
    """
    <div style='text-align: center; color: gray; font-size: 12px;'>
    シフト最適化システム v1.0 | OR-Tools & Streamlit
    </div>
    """,
    unsafe_allow_html=True
)

# 既存の結果ファイルを表示
if os.path.exists('shift_schedule_final.xlsx') and not st.session_state.get('just_created'):
    st.markdown("---")
    st.subheader("📂 以前の結果")
    try:
        df_existing = pd.read_excel('shift_schedule_final.xlsx', sheet_name='シフト表')
        with st.expander("以前のシフト表を表示"):
            # データ型を文字列に変換
            df_existing = df_existing.astype(str)
            st.dataframe(df_existing, use_container_width=True, height=300)
            with open('shift_schedule_final.xlsx', 'rb') as f:
                st.download_button(
                    label="⬇️ この結果をダウンロード",
                    data=f,
                    file_name="shift_schedule_final.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
    except:
        pass