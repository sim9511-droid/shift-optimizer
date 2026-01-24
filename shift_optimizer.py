import pandas as pd
from ortools.sat.python import cp_model
import datetime
from openpyxl.styles import Border, Side, Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import yaml

def load_constraints(yaml_file='constraints.yaml'):
    """YAMLファイルから制約条件を読み込む"""
    with open(yaml_file, 'r', encoding='utf-8') as f:
        config = yaml.safe_load(f)
    return config

def solve():
    input_file = 'Shift_Input.xlsx'
    
    # --- 1. データの読み込み ---
    # YAML設定を読み込み
    constraints = load_constraints()
    
    # 設定シート：1行目が「年」、2行目が「月」
    df_set = pd.read_excel(input_file, sheet_name='設定', header=None)
    year = int(df_set.iloc[0, 1])
    month = int(df_set.iloc[1, 1])
    
    # 休日回数シート
    df_staff = pd.read_excel(input_file, sheet_name='休日回数')
    
    # 希望休シート：A1削除により1行目がヘッダー(header=0)
    df_wish = pd.read_excel(input_file, sheet_name='希望休', header=0)
    # 1列目を日付型へ変換（空白などは無視）
    df_wish['日付'] = pd.to_datetime(df_wish.iloc[:, 0], errors='coerce').dt.date

    # シフト期間の計算
    period_config = constraints['period']
    start_day = period_config['start_day']
    end_day = period_config['end_day']
    
    # 祝日を設定シートのD列から読み込み
    holidays_list = []
    for i in range(len(df_set)):
        if pd.notna(df_set.iloc[i, 3]):  # D列（index=3）
            try:
                holiday_date = pd.to_datetime(df_set.iloc[i, 3])
                holidays_list.append(holiday_date.strftime('%Y-%m-%d'))
            except:
                pass
    
    # 21日〜翌月20日 (または指定された日付範囲)
    date_list = []
    start_date = datetime.date(year, month, start_day)
    if end_day < start_day:  # 翌月に跨ぐ場合
        if month == 12:
            end_date = datetime.date(year + 1, 1, end_day)
        else:
            end_date = datetime.date(year, month + 1, end_day)
    else:
        end_date = datetime.date(year, month, end_day)
    
    tmp = start_date
    while tmp <= end_date:
        date_list.append(tmp)
        tmp += datetime.timedelta(days=1)

    model = cp_model.CpModel()
    shifts = {}
    soba = {}

    # --- 2. モデル構築 (希望休の強制反映) ---
    for _, s_row in df_staff.iterrows():
        name = str(s_row['名前'])
        clean_name = name.replace('●', '').strip()
        
        for d_idx, d_date in enumerate(date_list):
            shifts[(name, d_idx)] = model.NewBoolVar(f's_{name}_{d_idx}')
            
            # 希望休・有給チェック
            is_off_requested = False
            if clean_name in df_wish.columns:
                # 該当日の行を抽出
                target_row = df_wish[df_wish['日付'] == d_date]
                if not target_row.empty:
                    val = target_row[clean_name].values[0]
                    # NaNでなく、「希」「有」「休」のいずれかが含まれていれば休み確定
                    if pd.notna(val):
                        val_str = str(val).strip()
                        if any(x in val_str for x in ['希', '有', '休']):
                            is_off_requested = True
            
            if is_off_requested:
                model.Add(shifts[(name, d_idx)] == 0)

            # そば打ち
            if s_row['そば打ち'] == '〇':
                soba[(name, d_idx)] = model.NewBoolVar(f'sb_{name}_{d_idx}')
                model.AddImplication(soba[(name, d_idx)], shifts[(name, d_idx)])
            else:
                soba[(name, d_idx)] = model.NewConstant(0)

    # --- 3. 制約条件 (人数枠など) ---
    rules = constraints['rules']
    max_consecutive = rules['max_consecutive_work_days']
    
    # Excelの設定シートから平日・休日の人数を読み込み
    weekday_min = int(df_set.iloc[3, 1])  # 平日最低人数
    weekday_max = int(df_set.iloc[4, 1])  # 平日最大人数
    holiday_min = int(df_set.iloc[5, 1])  # 休日最低人数
    holiday_max = int(df_set.iloc[6, 1])  # 休日最大人数
    
    # 祝日判定用
    holiday_dates = [datetime.datetime.strptime(h, '%Y-%m-%d').date() for h in holidays_list]
    
    for d_idx, d_date in enumerate(date_list):
        is_weekend = (d_date.weekday() >= 5)  # 土日
        is_holiday = is_weekend or (d_date in holiday_dates)  # 祝日も含める
        
        # 土日祝か平日かで最小・最大人数を切り替え
        min_staff = holiday_min if is_holiday else weekday_min
        max_staff = holiday_max if is_holiday else weekday_max
        
        # 最小人数制約
        model.Add(sum(shifts[(s['名前'], d_idx)] for _, s in df_staff.iterrows()) >= min_staff)
        
        # 最大人数制約
        model.Add(sum(shifts[(s['名前'], d_idx)] for _, s in df_staff.iterrows()) <= max_staff)
        
        # そば打ち1名制約
        model.Add(sum(soba[(s['名前'], d_idx)] for _, s in df_staff.iterrows()) == 1)

    for _, s_row in df_staff.iterrows():
        name = s_row['名前']
        # 月間合計出勤日数 = 全日数 - 月間休日数
        model.Add(sum(shifts[(name, d_idx)] for d_idx in range(len(date_list))) == (len(date_list) - int(s_row['月間休日数'])))
        # 連勤制限（YAMLの設定値を使用）
        for d_idx in range(len(date_list) - max_consecutive):
            model.Add(sum(shifts[(name, d_idx + k)] for k in range(max_consecutive + 1)) <= max_consecutive)
        
        # 連続休日を最大3日に制限（休みを分散させる）
        for d_idx in range(len(date_list) - 3):
            model.Add(sum(1 - shifts[(name, d_idx + k)] for k in range(4)) <= 3)
    
    # そば打ちを古藤・石橋・佐伯で公平に配分
    soba_staff = []
    for _, s_row in df_staff.iterrows():
        name = str(s_row['名前'])
        clean_name = name.replace('●', '').strip()
        if s_row['そば打ち'] == '〇' and clean_name in ['古藤', '石橋', '佐伯']:
            soba_staff.append(name)
    
    if len(soba_staff) >= 2:
        # 各スタッフのそば打ち回数を計算
        soba_counts = []
        for name in soba_staff:
            count = sum(soba[(name, d_idx)] for d_idx in range(len(date_list)))
            soba_counts.append(count)
        
        # 最大と最小の差を2回以内に制限（公平性を保つ）
        for i in range(len(soba_counts)):
            for j in range(i+1, len(soba_counts)):
                model.Add(soba_counts[i] - soba_counts[j] <= 2)
                model.Add(soba_counts[j] - soba_counts[i] <= 2)
    
    # そば打ちの連続を最小化（できる限り毎日交代する）
    # 連続する2日間で同じ人がそば打ちを担当しないようにする
    consecutive_soba_penalties = []
    for name in soba_staff:
        for d_idx in range(len(date_list) - 1):
            # 連続そば打ちの検出用変数
            consecutive = model.NewBoolVar(f'consecutive_soba_{name}_{d_idx}')
            # 2日連続でそば打ちをした場合にconsecutive=1
            model.Add(soba[(name, d_idx)] + soba[(name, d_idx + 1)] == 2).OnlyEnforceIf(consecutive)
            model.Add(soba[(name, d_idx)] + soba[(name, d_idx + 1)] <= 1).OnlyEnforceIf(consecutive.Not())
            consecutive_soba_penalties.append(consecutive)
    
    # 連続そば打ちを最小化する（目的関数）
    if consecutive_soba_penalties:
        model.Minimize(sum(consecutive_soba_penalties))

    # --- 4. 実行と結果出力 ---
    solver = cp_model.CpSolver()
    status = solver.Solve(model)

    if status in [cp_model.OPTIMAL, cp_model.FEASIBLE]:
        res = []
        for d_idx, d_date in enumerate(date_list):
            row = {'日付': f"{d_date.month}/{d_date.day}", '曜': "月火水木金土日"[d_date.weekday()]}
            for _, s_row in df_staff.iterrows():
                name = str(s_row['名前'])
                disp_name = name.replace('●', '').strip() # 出力から●を消す
                
                if solver.Value(shifts[(name, d_idx)]) == 1:
                    is_sb = (solver.Value(soba[(name, d_idx)]) == 1)
                    work_time = str(s_row['そば打ち時の時間' if is_sb else '基本勤務時間']).replace(':00','')
                    row[disp_name] = work_time + ("㋞" if is_sb else "")
                else:
                    # 希望休シートから元の表記（希休/有休）を探す
                    orig_val = ""
                    if disp_name in df_wish.columns:
                        m_row = df_wish[df_wish['日付'] == d_date]
                        if not m_row.empty:
                            orig_val = str(m_row[disp_name].values[0])
                    
                    if '有' in orig_val: row[disp_name] = '有休'
                    elif '希' in orig_val: row[disp_name] = '希休'
                    else: row[disp_name] = '休'
            res.append(row)

        # --- 計算結果を追加 ---
        # 各スタッフの統計情報を計算
        stats_rows = []
        
        # 公休数
        ko_kyuu_row = {'日付': '公休数', '曜': ''}
        for _, s_row in df_staff.iterrows():
            name = str(s_row['名前'])
            disp_name = name.replace('●', '').strip()
            ko_kyuu_count = 0
            for d_idx, d_date in enumerate(date_list):
                if solver.Value(shifts[(name, d_idx)]) == 0:
                    # 希望休シートから元の表記を確認
                    orig_val = ""
                    if disp_name in df_wish.columns:
                        m_row = df_wish[df_wish['日付'] == d_date]
                        if not m_row.empty:
                            orig_val = str(m_row[disp_name].values[0])
                    
                    # 有休でなければ公休
                    if '有' not in orig_val:
                        ko_kyuu_count += 1
            ko_kyuu_row[disp_name] = ko_kyuu_count
        stats_rows.append(ko_kyuu_row)
        
        # 有休数
        yu_kyuu_row = {'日付': '有休数', '曜': ''}
        for _, s_row in df_staff.iterrows():
            name = str(s_row['名前'])
            disp_name = name.replace('●', '').strip()
            yu_kyuu_count = 0
            for d_idx, d_date in enumerate(date_list):
                if solver.Value(shifts[(name, d_idx)]) == 0:
                    # 希望休シートから元の表記を確認
                    orig_val = ""
                    if disp_name in df_wish.columns:
                        m_row = df_wish[df_wish['日付'] == d_date]
                        if not m_row.empty:
                            orig_val = str(m_row[disp_name].values[0])
                    
                    # 有休カウント
                    if '有' in orig_val:
                        yu_kyuu_count += 1
            yu_kyuu_row[disp_name] = yu_kyuu_count
        stats_rows.append(yu_kyuu_row)
        
        # 総勤務時間（休憩時間を引く前）
        total_hours_row = {'日付': '総時間数', '曜': ''}
        staff_total_hours = {}  # 後で使うため保存
        for _, s_row in df_staff.iterrows():
            name = str(s_row['名前'])
            disp_name = name.replace('●', '').strip()
            total_hours = 0
            for d_idx in range(len(date_list)):
                if solver.Value(shifts[(name, d_idx)]) == 1:
                    is_sb = (solver.Value(soba[(name, d_idx)]) == 1)
                    work_time_str = str(s_row['そば打ち時の時間' if is_sb else '基本勤務時間'])
                    # 時間を計算 (HH:MM形式を想定)
                    try:
                        start_time, end_time = work_time_str.split('-')
                        start_h, start_m = map(int, start_time.split(':'))
                        end_h, end_m = map(int, end_time.split(':'))
                        hours = (end_h + end_m/60) - (start_h + start_m/60)
                        total_hours += hours
                    except:
                        total_hours += 9  # デフォルト9時間
            staff_total_hours[disp_name] = total_hours
            total_hours_row[disp_name] = int(total_hours) if total_hours == int(total_hours) else round(total_hours, 1)
        stats_rows.append(total_hours_row)
        
        # 休憩時間を引いた実働時間
        work_hours_row = {'日付': '実働時間', '曜': ''}
        staff_work_hours = {}  # 後で使うため保存
        for _, s_row in df_staff.iterrows():
            name = str(s_row['名前'])
            disp_name = name.replace('●', '').strip()
            total = staff_total_hours[disp_name]
            work_days = len(date_list) - int(s_row['月間休日数'])
            avg_hours_per_day = total / work_days if work_days > 0 else 0
            
            # 休憩時間の計算
            if avg_hours_per_day < 6:
                break_time = 0
            elif avg_hours_per_day <= 8:
                break_time = 0.75 * work_days  # 45分 = 0.75時間
            else:
                break_time = 1.0 * work_days  # 1時間
            
            work_hours = total - break_time
            staff_work_hours[disp_name] = work_hours
            work_hours_row[disp_name] = int(work_hours) if work_hours == int(work_hours) else round(work_hours, 1)
        stats_rows.append(work_hours_row)
        
        # 法定時間数（月の日数×40÷7）
        days_in_period = len(date_list)
        legal_hours = days_in_period * 40 / 7
        legal_hours_row = {'日付': '法定時間数', '曜': ''}
        for _, s_row in df_staff.iterrows():
            name = str(s_row['名前'])
            disp_name = name.replace('●', '').strip()
            legal_hours_row[disp_name] = round(legal_hours, 1)
        stats_rows.append(legal_hours_row)
        
        # 差分（実働時間 - 法定時間数）
        diff_hours_row = {'日付': '差分', '曜': ''}
        for _, s_row in df_staff.iterrows():
            name = str(s_row['名前'])
            disp_name = name.replace('●', '').strip()
            diff = staff_work_hours[disp_name] - legal_hours
            diff_hours_row[disp_name] = f"{diff:+.1f}"  # プラス記号を付ける
        stats_rows.append(diff_hours_row)
        
        # --- 出勤人数を追加 ---
        for i, row in enumerate(res):
            count = 0
            for _, s_row in df_staff.iterrows():
                name = str(s_row['名前'])
                disp_name = name.replace('●', '').strip()
                if disp_name in row and row[disp_name] != '休' and row[disp_name] != '希休' and row[disp_name] != '有休':
                    count += 1
            row['人'] = count
        
        # stats_rowsにも人列を追加
        for stats_row in stats_rows:
            stats_row['人'] = ''
        
        df_out = pd.DataFrame(res + stats_rows)
        with pd.ExcelWriter('shift_schedule_final.xlsx', engine='openpyxl') as writer:
            df_out.to_excel(writer, index=False, sheet_name='シフト表')
            ws = writer.sheets['シフト表']
            
            # 祝日リストをdate型に変換
            holiday_dates = [datetime.datetime.strptime(h, '%Y-%m-%d').date() for h in holidays_list]
            
            # 列の幅を調整
            ws.column_dimensions['A'].width = 10  # 日付
            ws.column_dimensions['B'].width = 6   # 曜
            
            # スタッフ列の幅
            for col_idx, col_letter in enumerate([get_column_letter(i) for i in range(3, len(df_out.columns) + 1)], start=3):
                ws.column_dimensions[col_letter].width = 12
            
            # セルスタイルを適用
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                               top=Side(style='thin'), bottom=Side(style='thin'))
            center_align = Alignment(horizontal='center', vertical='center')
            
            # 色定義
            blue_font = Font(color='0070C0', size=11)  # 青文字
            red_font = Font(color='FF0000', size=11)   # 赤文字
            black_font = Font(color='000000', size=11)  # 黒文字
            holiday_fill = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')  # 背景赤
            holiday_font = Font(color='FFFFFF', bold=True, size=11)  # 白文字、太字
            kibo_fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')  # 背景黄色
            yukyuu_fill = PatternFill(start_color='FFA500', end_color='FFA500', fill_type='solid')  # 背景オレンジ
            kyuu_fill = PatternFill(start_color='D3D3D3', end_color='D3D3D3', fill_type='solid')  # 背景薄い灰色
            
            # ヘッダー行（1行目）
            for col_idx, cell in enumerate(ws[1], start=1):
                cell.border = thin_border
                cell.alignment = center_align
                cell.font = Font(bold=True)
            
            # データ行（シフト表の日付行のみ、統計行は除く）
            for row_idx in range(2, len(res) + 2):  # 2行目から日付行の最後まで
                res_item = res[row_idx - 2]
                for col_idx, cell in enumerate(ws[row_idx], start=1):
                    cell.border = thin_border
                    cell.alignment = center_align
                    
                    # 日付列
                    if col_idx == 1:
                        cell.font = Font(size=11)
                    # 曜列（色付けはここだけ）
                    elif col_idx == 2:
                        if cell.value and len(str(cell.value)) > 0:
                            weekday_char = str(cell.value)
                            # 祝日判定
                            date_str = res_item.get('日付', '')
                            is_holiday_date = False
                            if date_str and '/' in date_str:
                                try:
                                    m, d = map(int, date_str.split('/'))
                                    check_date = datetime.date(year, m, d)
                                    is_holiday_date = check_date in holiday_dates
                                except:
                                    pass
                            
                            if is_holiday_date:
                                cell.fill = holiday_fill
                                cell.font = holiday_font
                            elif weekday_char == '土':
                                cell.font = blue_font
                            elif weekday_char == '日':
                                cell.font = red_font
                            else:
                                cell.font = Font(size=11)
                    # スタッフ列（黒文字で色分けなし、ただし背景色は付ける）
                    else:
                        cell_value = str(cell.value) if cell.value else ''
                        
                        if '希休' in cell_value:
                            cell.fill = kibo_fill
                            cell.font = black_font
                        elif '有休' in cell_value:
                            cell.fill = yukyuu_fill
                            cell.font = black_font
                        elif cell_value == '休':
                            cell.fill = kyuu_fill
                            cell.font = black_font
                        else:
                            cell.font = black_font
            
            # 統計行（色付けなし、罫線と中央揃えのみ）
            for row_idx in range(len(res) + 2, len(res) + len(stats_rows) + 2):
                for cell in ws[row_idx]:
                    cell.border = thin_border
                    cell.alignment = center_align
                    cell.font = Font(size=11)
        
        print("✅ 完了！色付けと出勤人数を追加した『shift_schedule_final.xlsx』を作成しました。")
    else:
        print("❌ 条件に合うシフトが見つかりませんでした。希望休が多すぎるか、人数設定が厳しい可能性があります。")

if __name__ == "__main__":
    solve()